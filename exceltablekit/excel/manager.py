from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from exceltablekit.errors import (
    ExcelSheetMissingError,
    HeaderTableNotDefinedError,
    InvalidCellError,
    InvalidColumnValueError,
    InvalidRowValueError,
    TableNotDefinedError,
)
from exceltablekit.table import Table
from exceltablekit.utils import Validation


class ExcelManager:
    _filepath: str
    _workbook: Workbook
    _worksheet: Worksheet
    _table: Table | None

    def __init__(
        self,
        filepath: str,
        create_sheet: bool = False,
        sheetname: str | None = None,
    ) -> None:
        self._filepath = filepath
        is_new_file = not Path(filepath).exists()
        self._workbook = Workbook() if is_new_file else load_workbook(filepath)
        self._worksheet = self._resolve_worksheet(sheetname, create_sheet, is_new_file)
        self._table = None

    def _resolve_worksheet(
        self, sheetname: str | None, create_sheet: bool, is_new_file: bool
    ) -> Worksheet:
        """Resolves which worksheet to use based on initialization parameters."""
        if not sheetname:
            active = self._workbook.active
            if active is None:
                raise ExcelSheetMissingError("No active sheet found in the workbook")
            return active

        if create_sheet or is_new_file:
            return self._workbook.create_sheet(sheetname)

        if not self._is_sheet_exists(sheetname):
            raise ExcelSheetMissingError(sheetname)
        return self._workbook[sheetname]

    def _is_sheet_exists(self, sheetname: str) -> bool:
        """Returns true if a specific worksheet exists within the workbook.

        Parameters:
        -----------
        sheetname : str
            Sheet name of worksheet.
        """
        sheets = self._workbook.sheetnames
        return sheetname in sheets

    def _get_column(self, cell: str) -> str:
        """Gets the column part of a cell address."""
        return "".join([character for character in cell if character.isalpha()])

    def _get_row(self, cell: str) -> int:
        """Gets the row part of a cell address."""
        return int("".join([character for character in cell if character.isdigit()]))

    def set_table(self, start_cell: str, end_cell: str) -> None:
        """Defines the table area within the Excel worksheet.

        Parameters:
        ----------
        start_cell : str
            Starting cell of the table.
        end_cell : str
            Ending cell of the table.

        Example:
        -------
        >> my_excel = ExcelManager("my_excel.xlsx")
        >> my_excel.set_table("A1", "D4")
        """
        start_cell = start_cell.upper()
        end_cell = end_cell.upper()
        try:
            Validation.cell(start_cell)
            Validation.cell(end_cell)
        except (InvalidCellError, InvalidColumnValueError, InvalidRowValueError):
            raise

        start_column = self._get_column(start_cell)
        end_column = self._get_column(end_cell)
        start_row = self._get_row(start_cell)
        end_row = self._get_row(end_cell)
        self._table = Table(
            worksheet=self._worksheet,
            first_column=start_column,
            first_row=start_row,
            last_column=end_column,
            last_row=end_row,
        )

    def define_header(self, rows: int) -> None:
        """Defines the number of rows that the table will use as its header.

        Parameters:
        ----------
        rows : int
            Total count of rows designated as headers.
        """
        if self._table:
            self._table.set_header(rows)

    def apply_style_header(self, property: str, value: Any) -> None:
        """Applies a style to the header Excel table.

        Parameters:
        ----------
        property : str
            Property to define.
        value : Any
            Value of the property to set.
        """
        if not self._table:
            raise TableNotDefinedError()
        if not self._table.get_header():
            raise HeaderTableNotDefinedError()
        self._table.set_style_header(property, value)
        self._workbook.save(self._filepath)

    def apply_style_body(self, property: str, value: Any) -> None:
        """Applies a style to the body Excel table.

        Parameters:
        ----------
        property : str
            Property to define.
        value : Any
            Value of the property to set.
        """
        if not self._table:
            raise TableNotDefinedError()
        self._table.set_style_body(property, value)
        self._workbook.save(self._filepath)

    def freeze_header(self) -> None:
        """Freezes the header rows of the table so they remain visible when scrolling.

        Example:
        -------
        >> my_excel = ExcelManager("my_excel.xlsx")
        >> my_excel.set_table("A1", "D10")
        >> my_excel.define_header(rows=2)
        >> my_excel.freeze_header()
        """
        if not self._table:
            raise TableNotDefinedError()
        if not self._table.get_header():
            raise HeaderTableNotDefinedError()
        first_column = self._table.get_first_column()
        freeze_row = self._table.get_first_row() + self._table.get_total_row_headers()
        self._worksheet.freeze_panes = f"{first_column}{freeze_row}"
        self._workbook.save(self._filepath)

    def auto_fit_columns(self) -> None:
        """Adjusts each column width to fit the longest cell content in the table.

        Example:
        -------
        >> my_excel = ExcelManager("my_excel.xlsx")
        >> my_excel.set_table("A1", "D10")
        >> my_excel.auto_fit_columns()
        """
        if not self._table:
            raise TableNotDefinedError()
        all_cells = self._table.get_header() + self._table.get_body()
        col_widths: dict[str, int] = {}
        for cell in all_cells:
            col = cell.get_column()
            ws_cell = self._worksheet[cell.get_cell()]
            value_len = len(str(ws_cell.value)) if ws_cell.value is not None else 0
            col_widths[col] = max(col_widths.get(col, 0), value_len)
        for col, width in col_widths.items():
            self._worksheet.column_dimensions[col].width = width + 2
        self._workbook.save(self._filepath)

    def apply_style_table(self, property: str, value: Any) -> None:
        """Applies a style to the Excel table.

        Parameters:
        ----------
        property : str
            Property to define.
        value : Any
            Value of the property to set.
        """
        if not self._table:
            raise TableNotDefinedError()
        if self._table.get_header():
            self._table.set_style_header(property, value)
        self._table.set_style_body(property, value)
        self._workbook.save(self._filepath)
