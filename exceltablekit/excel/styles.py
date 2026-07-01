from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color

from exceltablekit.constants import BorderStyle
from exceltablekit.errors import (
    HeaderTableAlignmentError,
    HeaderTableBackgroundError,
    HeaderTableBorderError,
    HeaderTableFontError,
    HeaderTableFreezeError,
    HeaderTableNotDefinedError,
    TableAlignmentError,
    TableBackgroundError,
    TableBorderError,
    TableColumnWidthError,
    TableFontError,
    TableNotDefinedError,
)
from exceltablekit.excel.manager import ExcelManager


class ExcelStyle:
    _excel: ExcelManager

    def __init__(self, excel: ExcelManager) -> None:
        self._excel = excel

    def set_background(self, hex_color: str) -> None:
        """Sets the fill color for all cells in the Excel table.

        Parameters:
        ----------
        hex_color : str
            Hexadecimal color to apply.
        """
        try:
            color = Color(rgb=hex_color, type="rgb")
            fill = PatternFill(patternType="solid", start_color=color, end_color=color)
            self._excel.apply_style_body(property="fill", value=fill)
        except TableNotDefinedError:
            raise
        except Exception as error:
            raise TableBackgroundError(error) from error

    def set_font(
        self,
        type: str = "Time New Roman",
        size: float = 10.0,
        bold: bool = False,
        italic: bool = False,
        hex_color: str = "FFFF0000",
    ) -> None:
        """Sets the font for all cells in the Excel table.

        Parameters:
        ----------
        type : str, optional
            Font type name.
        size : float, optional
            Size font.
        bold : bool, optional
            Bold type.
        italic : bool, optional
            Italic type.
        hex_color : str, optional
            Hexadecimal color to apply.
        """
        try:
            color = Color(rgb=hex_color, type="rgb")
            font = Font(name=type, size=size, bold=bold, italic=italic, color=color)
            self._excel.apply_style_body(property="font", value=font)
        except TableNotDefinedError:
            raise
        except Exception as error:
            raise TableFontError(error) from error

    def set_alignment(
        self,
        horizontal: str = "left",
        vertical: str = "center",
        wrap_text: bool = False,
    ) -> None:
        """Sets the alignment text for all cells in the Excel table.

        Parameters:
        ----------
        horizontal : str, optional
            Horizontal text centering.
        vertical : str, optional
            Vertical text centering.
        wrap_text : bool, optional
            Wrap text.
        """
        try:
            alignment = Alignment(
                horizontal=horizontal, vertical=vertical, wrap_text=wrap_text
            )
            self._excel.apply_style_body(property="alignment", value=alignment)
        except TableNotDefinedError:
            raise
        except Exception as error:
            raise TableAlignmentError(error) from error

    def set_border(
        self,
        left: bool = True,
        left_style: BorderStyle = "thin",
        right: bool = True,
        right_style: BorderStyle = "thin",
        top: bool = True,
        top_style: BorderStyle = "thin",
        bottom: bool = True,
        bottom_style: BorderStyle = "thin",
        hex_color: str = "000000",
    ) -> None:
        """Sets the border for all cells in the Excel table.

        Parameters:
        ----------
        left : bool, optional
            If True, enables the left border.
        left_style : BorderStyle, optional
            The style or type of the left border.
        right : bool, optional
            If True, enables the right border.
        right_style : BorderStyle, optional
            The style or type of the right border.
        top : bool, optional
            If True, enables the top border.
        top_style : BorderStyle, optional
            The style or type of the top border.
        bottom : bool, optional
            If True, enables the bottom border.
        bottom_style : BorderStyle, optional
            The style or type of the bottom border.
        hex_color : str, optional
            Hexadecimal color to apply.
        """
        try:
            color = Color(rgb=hex_color, type="rgb")
            border = Border(
                left=Side(style=left_style, color=color) if left else None,
                right=Side(style=right_style, color=color) if right else None,
                top=Side(style=top_style, color=color) if top else None,
                bottom=Side(style=bottom_style, color=color) if bottom else None,
            )
            self._excel.apply_style_body(property="border", value=border)
        except TableNotDefinedError:
            raise
        except Exception as error:
            raise TableBorderError(error) from error

    def auto_fit_columns(self) -> None:
        """Adjusts each column width to fit the longest cell content in the table.

        Example:
        -------
        >>> styles = ExcelStyle(excel)
        >>> styles.auto_fit_columns()
        """
        try:
            self._excel.auto_fit_columns()
        except TableNotDefinedError:
            raise
        except Exception as error:
            raise TableColumnWidthError(error) from error

    # ----------------- HEADERS ---------------------------------

    def set_header_background(self, hex_color: str) -> None:
        """Sets the fill color for all cells in the header Excel table.

        Parameters:
        ----------
        hex_color : str
            Hexadecimal color to apply.
        """
        try:
            color = Color(rgb=hex_color, type="rgb")
            fill = PatternFill(patternType="solid", start_color=color, end_color=color)
            self._excel.apply_style_header(property="fill", value=fill)
        except (TableNotDefinedError, HeaderTableNotDefinedError):
            raise
        except Exception as error:
            raise HeaderTableBackgroundError(error) from error

    def set_header_font(
        self,
        type: str = "Time New Roman",
        size: float = 10.0,
        bold: bool = False,
        italic: bool = False,
        hex_color: str = "FFFF0000",
    ) -> None:
        """Sets the font for all cells in the header Excel table.

        Parameters:
        ----------
        type : str, optional
            Font type name.
        size : float, optional
            Size font.
        bold : bool, optional
            Bold type.
        italic : bool, optional
            Italic type.
        hex_color : str, optional
            Hexadecimal color to apply.
        """
        try:
            color = Color(rgb=hex_color, type="rgb")
            font = Font(name=type, size=size, bold=bold, italic=italic, color=color)
            self._excel.apply_style_header(property="font", value=font)
        except (TableNotDefinedError, HeaderTableNotDefinedError):
            raise
        except Exception as error:
            raise HeaderTableFontError(error) from error

    def set_header_alignment(
        self,
        horizontal: str = "left",
        vertical: str = "center",
        wrap_text: bool = False,
    ) -> None:
        """Sets the alignment text for all cells in the header Excel table.

        Parameters:
        ----------
        horizontal : str, optional
            Horizontal text centering.
        vertical : str, optional
            Vertical text centering.
        wrap_text : bool, optional
            Wrap text.
        """
        try:
            alignment = Alignment(
                horizontal=horizontal, vertical=vertical, wrap_text=wrap_text
            )
            self._excel.apply_style_header(property="alignment", value=alignment)
        except (TableNotDefinedError, HeaderTableNotDefinedError):
            raise
        except Exception as error:
            raise HeaderTableAlignmentError(error) from error

    def freeze_header(self) -> None:
        """Freezes the header rows so they remain fixed when scrolling in Excel.

        Example:
        -------
        >>> styles = ExcelStyle(excel)
        >>> styles.freeze_header()
        """
        try:
            self._excel.freeze_header()
        except (TableNotDefinedError, HeaderTableNotDefinedError):
            raise
        except Exception as error:
            raise HeaderTableFreezeError(error) from error

    def set_header_border(
        self,
        left: bool = True,
        left_style: BorderStyle = "thin",
        right: bool = True,
        right_style: BorderStyle = "thin",
        top: bool = True,
        top_style: BorderStyle = "thin",
        bottom: bool = True,
        bottom_style: BorderStyle = "thin",
        hex_color: str = "000000",
    ) -> None:
        """Sets the border for all cells in the header Excel table.

        Parameters:
        ----------
        left : bool, optional
            If True, enables the left border.
        left_style : BorderStyle, optional
            The style or type of the left border.
        right : bool, optional
            If True, enables the right border.
        right_style : BorderStyle, optional
            The style or type of the right border.
        top : bool, optional
            If True, enables the top border.
        top_style : BorderStyle, optional
            The style or type of the top border.
        bottom : bool, optional
            If True, enables the bottom border.
        bottom_style : BorderStyle, optional
            The style or type of the bottom border.
        hex_color : str, optional
            Hexadecimal color to apply.
        """
        try:
            color = Color(rgb=hex_color, type="rgb")
            border = Border(
                left=Side(style=left_style, color=color) if left else None,
                right=Side(style=right_style, color=color) if right else None,
                top=Side(style=top_style, color=color) if top else None,
                bottom=Side(style=bottom_style, color=color) if bottom else None,
            )
            self._excel.apply_style_header(property="border", value=border)
        except (TableNotDefinedError, HeaderTableNotDefinedError):
            raise
        except Exception as error:
            raise HeaderTableBorderError(error) from error
